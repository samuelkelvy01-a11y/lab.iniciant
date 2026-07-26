from openpyxl import Workbook, load_workbook
import os

OUT = "outputs"
os.makedirs(OUT, exist_ok=True)

xlsx_path = os.path.join(OUT, "notas.xlsx")

# Criar uma planilha nova
wb = Workbook()
ws = wb.active
ws.title = "Notas"

ws["A1"] = "Nome"
ws["B1"] = "Nota"

ws.append(["Samuel", 9])
ws.append(["João", 7])
ws.append(["Maria", 8.5])

wb.save(xlsx_path)
print(f"Arquivo salvo em: {xlsx_path}")

# Ler e mostrar conteúdo
wb2 = load_workbook(xlsx_path)
ws2 = wb2["Notas"]
print("\nConteúdo da planilha:")
for row in ws2.iter_rows(values_only=True):
    print(row)
